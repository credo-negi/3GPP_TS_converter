import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

from tests import support  # noqa: F401
from ts_converter.ir import (Cell, DisplayMath, Document, Para, Section,
                             Seg, Table)
from ts_converter.md_writer import ImageStore, MdWriter
from ts_converter.xlsx_writer import write_xlsx


def t(s):
    return Seg("text", s)


def make_doc() -> Document:
    doc = Document("38211", 15, "15.0.0", "t", source="x.docx")
    doc.sections.append(Section(0, "", "Cover", [Para("cover", [t("TS")])]))
    s1 = Section(1, "5", "Generic", [Para("text", [t("Intro. Two.")])])
    s2 = Section(2, "5.1", "Mapper", [
        Para("text", [t("Value "), Seg("math", "x_{1}"),
                      t(" is used. Next one.")]),
        Para("bullet", [t("first. Second.")], 1),
        Para("bullet", [t("deep")], 2),
        Para("note", [t("NOTE 1: be careful. Really.")]),
        DisplayMath(r"a=b"),
        Table([[Cell([[t("A")]]), Cell([[t("B")]], colspan=2),
                Cell(covered=True, origin=(0, 1))],
               [Cell([[t("1")]], rowspan=2), Cell([[t("x | y")]]),
                Cell([[t("z")]])],
               [Cell(covered=True, origin=(1, 0)), Cell([[t("q")]]),
                Cell([[Seg("math", r"x_{1}<2")]])]]),
        Table([[Cell([[t("H1")]]), Cell([[t("H2")]])],
               [Cell([[t("a | b")]]), Cell([[t("c")]])],
               [Cell([[Seg("math", "y")]]), Cell([[t("c")]])]]),
    ])
    doc.sections += [s1, s2]
    return doc


class MdWriterTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name)
        zpath = self.out / "d.docx"
        with zipfile.ZipFile(zpath, "w"):
            pass
        self.zf = zipfile.ZipFile(zpath)
        doc = make_doc()
        self.w = MdWriter(doc, ImageStore(self.zf, self.out, "r"),
                          self.out / "md")
        self.w.write()

    def tearDown(self):
        self.zf.close()
        self.tmp.cleanup()

    def read(self, name):
        return (self.out / "md" / name).read_text("utf-8")

    def test_files_per_clause(self):
        names = sorted(p.name for p in (self.out / "md").glob("*.md"))
        self.assertEqual(names, ["000_cover.md", "001_5_Generic.md",
                                 "002_5.1_Mapper.md", "index.md"])

    def test_one_sentence_per_line(self):
        txt = self.read("001_5_Generic.md")
        self.assertIn("Intro.\nTwo.\n", txt)

    def test_inline_math_and_bullets(self):
        txt = self.read("002_5.1_Mapper.md")
        self.assertIn("Value $x_{1}$ is used.\nNext one.", txt)
        self.assertIn("- first.\n  Second.\n  - deep", txt)
        self.assertIn("> NOTE 1: be careful.\n> Really.", txt)
        self.assertIn("$$a=b$$", txt)

    def test_table_merged_cells_html(self):
        txt = self.read("002_5.1_Mapper.md")
        self.assertIn("<table>", txt)
        self.assertIn('<th>A</th>', txt)
        self.assertIn('<th colspan="2">B</th>', txt)
        self.assertIn('<td rowspan="2">1</td>', txt)
        self.assertIn("<td>x | y</td>", txt)
        self.assertEqual(txt.count("<td>q</td>"), 1)
        self.assertEqual(txt.count("<tr>"), 3)

    def test_math_in_merged_table_is_mathml(self):
        txt = self.read("002_5.1_Mapper.md")
        self.assertIn("<td><math><msub><mi>x</mi><mn>1</mn></msub>"
                      "<mo>&lt;</mo><mn>2</mn></math></td>", txt)
        table = txt[txt.index("<table>"):txt.index("</table>")]
        self.assertNotIn("$", table)
        self.assertNotIn("\n\n", table)      # one markdown block

    def test_unconvertible_math_in_html_table_stays_latex(self):
        from ts_converter.render import seg_atom
        a = seg_atom(Seg("math", r"\nosuchcommand<1"), "html", str)
        self.assertEqual(a, r"$\nosuchcommand&lt;1$")

    def test_math_in_plain_table_stays_latex(self):
        txt = self.read("002_5.1_Mapper.md")
        self.assertIn("| $y$ | c |", txt)

    def test_table_without_merge_is_gfm(self):
        txt = self.read("002_5.1_Mapper.md")
        self.assertIn("| H1 | H2 |", txt)
        self.assertIn("| a \\| b | c |", txt)

    def test_breadcrumb_and_index(self):
        self.assertTrue(self.read("002_5.1_Mapper.md").startswith(
            "<!-- TS 38.211 Rel-15 V15.0.0 | 5 Generic > 5.1 Mapper -->"))
        self.assertIn("  - [5.1 Mapper](002_5.1_Mapper.md)",
                      self.read("index.md"))


class XlsxWriterTest(unittest.TestCase):
    def test_cells(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "a.xlsx"
            write_xlsx(make_doc(), p)
            wb = openpyxl.load_workbook(p)
            self.assertEqual(wb.sheetnames, ["Index", "Cover", "5 Generic"])
            ws = wb["5 Generic"]
            col = [c.value for c in ws["C"]]
            self.assertIn("Intro.", col)
            self.assertIn("Two.", col)           # one sentence per cell
            self.assertIn("$a=b$", col)
            merged = {str(r) for r in ws.merged_cells.ranges}
            self.assertTrue(any(m.startswith("D") for m in merged))
            self.assertTrue(any(m.startswith("C") and m.endswith(
                str(int(m.split(":")[1][1:]))) for m in merged))


if __name__ == "__main__":
    unittest.main()
