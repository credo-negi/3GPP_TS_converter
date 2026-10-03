"""Document -> xlsx: one sentence per cell.

Layout (one sheet per top-level clause, plus an ``Index`` sheet):
  A: clause number, B: kind, C...: content
  text paragraphs : one sentence per row, in column C
  display math    : kind 'eq', LaTeX in column C
  tables          : one grid cell per table cell, from column C
                    (merged cells are really merged)
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .ir import Document, DisplayMath, ImageBlock, Para, Section, Table
from .render import segs_to_sentences

BAD_SHEET = re.compile(r"[\[\]\*\?/\\:]")
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
KIND_FILL = PatternFill("solid", fgColor="F2F2F2")
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sheet_name(sec: Section, used: set[str]) -> str:
    base = BAD_SHEET.sub(" ", (sec.number + " " + sec.title).strip())
    base = base[:28].strip() or "sheet"
    name, i = base, 2
    while name.lower() in used:
        name = f"{base[:25]}_{i}"
        i += 1
    used.add(name.lower())
    return name


def clean(s: str) -> str:
    return ILLEGAL.sub("", s)


def set_text(cell, s: str):
    """Always store text (a leading '=' must not become a formula)."""
    cell.value = clean(s)
    cell.data_type = "s"
    return cell


class Sheet:
    def __init__(self, ws):
        self.ws, self.row = ws, 1
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12

    def put(self, num: str, kind: str, values: list[str], bold=False,
            wrap=False):
        ws = self.ws
        ws.cell(self.row, 1, num)
        c = ws.cell(self.row, 2, kind)
        c.fill = KIND_FILL
        for i, v in enumerate(values):
            cell = set_text(ws.cell(self.row, 3 + i), v)
            if bold:
                cell.font = Font(bold=True)
            if wrap:
                cell.alignment = Alignment(wrap_text=True,
                                           vertical="top")
        self.row += 1


def write_xlsx(doc: Document, path: Path, img_path=lambda s: s):
    wb = Workbook()
    index = wb.active
    index.title = "Index"
    index.append(["TS", f"{doc.spec[:2]}.{doc.spec[2:]}"])
    index.append(["Release", doc.release])
    index.append(["Version", doc.version])
    index.append(["Source", doc.source])
    index.append([])
    index.append(["Sheet", "Clause", "Title"])
    used: set[str] = {"index"}
    cur: Sheet | None = None
    for sec in doc.sections:
        if sec.level <= 1:
            name = sheet_name(sec, used) if sec.level else "Cover"
            if sec.level == 0:
                used.add("cover")
            ws = wb.create_sheet(name)
            cur = Sheet(ws)
            index.append([name, sec.number, sec.title])
            ix = index.cell(index.max_row, 1)
            ix.hyperlink = f"#'{name}'!A1"
            ix.font = Font(color="0563C1", underline="single")
        assert cur is not None
        if sec.level > 0:
            cur.put(sec.number, f"H{sec.level}",
                    [sec.title], bold=True)
        for b in sec.blocks:
            write_block(cur, sec, b, img_path)
    for ws in wb.worksheets:
        if ws.title != "Index":
            ws.column_dimensions["C"].width = 60
            ws.freeze_panes = "C2"
    index.column_dimensions["A"].width = 32
    index.column_dimensions["C"].width = 70
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_block(sh: Sheet, sec: Section, b, img_path):
    num = sec.number
    if isinstance(b, Para):
        kind = {"bullet": f"bullet{b.level}", "note": "note",
                "example": "example", "def": "def", "caption": "caption",
                "cover": "cover"}.get(b.kind, "text")
        segs = b.segs
        if b.kind == "def":
            from .md_writer import MdWriter
            segs = MdWriter.split_tab(segs)
        for s in segs_to_sentences(segs, "xlsx", img_path):
            sh.put(num, kind, [s])
    elif isinstance(b, DisplayMath):
        sh.put(num, "eq", [f"${b.latex}$"])
    elif isinstance(b, ImageBlock):
        sh.put(num, "image", [img_path(b.src)])
    elif isinstance(b, Table):
        write_table(sh, num, b, img_path)


def write_table(sh: Sheet, num: str, t: Table, img_path):
    if not t.rows:
        return
    top = sh.row
    ncols = max(len(r) for r in t.rows)
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row):
            if c.covered:
                continue
            sents: list[str] = []
            for segs in c.paras:
                sents.extend(segs_to_sentences(segs, "xlsx", img_path))
            cell = set_text(sh.ws.cell(top + ri, 3 + ci), "\n".join(sents))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if ri == 0:
                cell.font = Font(bold=True)
                cell.fill = HEAD_FILL
            if c.colspan > 1 or c.rowspan > 1:
                sh.ws.merge_cells(
                    start_row=top + ri, start_column=3 + ci,
                    end_row=top + ri + c.rowspan - 1,
                    end_column=3 + ci + c.colspan - 1)
        sh.ws.cell(top + ri, 1, num)
        sh.ws.cell(top + ri, 2, "table").fill = KIND_FILL
    for ci in range(ncols):
        letter = get_column_letter(3 + ci)
        if ci > 0:
            cur = sh.ws.column_dimensions[letter].width or 0
            sh.ws.column_dimensions[letter].width = max(cur, 18)
    sh.row = top + len(t.rows)
